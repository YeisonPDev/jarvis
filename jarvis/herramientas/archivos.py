"""Herramientas de archivos: solo lectura y apertura en el programa predeterminado."""

from __future__ import annotations

import difflib
import os
import time

from jarvis.herramientas.base import (
    Contexto, Herramienta, Resultado, error, esquema_objeto, exito,
)
from jarvis.seguridad.errores import ErrorSeguridad
from jarvis.seguridad.lista_negra import es_ejecutable, normalizar_texto
from jarvis.seguridad.rutas import es_enlace

EXTENSIONES_TEXTO = {".txt", ".md", ".csv", ".json", ".log"}
EXTENSIONES_DOCUMENTO = {".docx", ".xlsx", ".pdf"}
EXTENSIONES_LEIBLES = EXTENSIONES_TEXTO | EXTENSIONES_DOCUMENTO

# Carpetas que no aportan nada y vuelven lenta la búsqueda.
_CARPETAS_IGNORADAS = {
    "node_modules", ".git", "__pycache__", ".venv", "venv", "$recycle.bin",
    "system volume information", ".cache", ".idea", ".vs",
}


def _tamano_legible(n: int) -> str:
    for unidad in ("B", "KB", "MB", "GB"):
        if n < 1024:
            return f"{n:.0f} {unidad}" if unidad == "B" else f"{n:.1f} {unidad}"
        n /= 1024
    return f"{n:.1f} TB"


# ---------------------------------------------------------------------- #
def listar_carpeta(ctx: Contexto, args: dict) -> Resultado:
    carpeta = ctx.validador.validar(args["ruta"], tipo="carpeta")
    limite = int(ctx.config.limites.max_elementos_listado)
    carpetas, archivos, ocultos = [], [], 0
    with os.scandir(carpeta) as it:
        for entrada in it:
            if ctx.validador.lista_negra.motivo_nombre(entrada.name):
                ocultos += 1
                continue
            try:
                if entrada.is_dir(follow_symlinks=False):
                    carpetas.append(entrada.name + "\\")
                else:
                    tam = entrada.stat(follow_symlinks=False).st_size
                    archivos.append(f"{entrada.name} ({_tamano_legible(tam)})")
            except OSError:
                continue
    elementos = sorted(carpetas, key=str.lower) + sorted(archivos, key=str.lower)
    total = len(elementos)
    texto = "\n".join(elementos[:limite]) or "(carpeta vacía)"
    resumen = f"Carpeta {carpeta}: {len(carpetas)} carpetas y {len(archivos)} archivos."
    if total > limite:
        resumen += f" Se muestran los primeros {limite}."
    if ocultos:
        resumen += f" {ocultos} elementos protegidos no se muestran."
    return exito(resumen, texto, origen=carpeta)


# ---------------------------------------------------------------------- #
def _puntaje(consulta: str, nombre: str) -> float:
    n = normalizar_texto(nombre)
    base = os.path.splitext(n)[0]
    if consulta == n or consulta == base:
        return 2.0
    if consulta in n:
        return 1.5
    palabras = [p for p in consulta.replace("_", " ").replace("-", " ").split() if p]
    if palabras and all(p in n for p in palabras):
        return 1.2
    return difflib.SequenceMatcher(None, consulta, base).ratio()


def buscar_archivos(ctx: Contexto, args: dict) -> Resultado:
    consulta = normalizar_texto(args["nombre_aproximado"].strip())
    if len(consulta) < 2:
        return error("El nombre a buscar debe tener al menos 2 caracteres.")
    lim = ctx.config.limites

    if args.get("carpeta_opcional"):
        raices = [ctx.validador.validar(args["carpeta_opcional"], tipo="carpeta")]
    else:
        raices = []
        for p in ctx.validador.permitidas_originales:
            try:
                raices.append(ctx.validador.validar(p, tipo="carpeta"))
            except ErrorSeguridad:
                continue  # carpeta permitida que no existe en este equipo

    fin = time.monotonic() + float(lim.tiempo_busqueda_s)
    revisados, encontrados, interrumpido = 0, [], False
    vistos: set[str] = set()

    for raiz in raices:
        for actual, subcarpetas, archivos in os.walk(raiz, followlinks=False):
            # Podar: carpetas protegidas, bloqueadas, enlaces/junctions y basura conocida.
            conservar = []
            for s in subcarpetas:
                ruta_s = os.path.join(actual, s)
                if (
                    normalizar_texto(s) in _CARPETAS_IGNORADAS
                    or ctx.validador.lista_negra.motivo_nombre(s)
                    or es_enlace(ruta_s)
                    or ctx.validador.esta_bloqueada(ruta_s)
                ):
                    continue
                conservar.append(s)
            subcarpetas[:] = conservar

            for nombre in archivos + [s for s in subcarpetas]:
                revisados += 1
                ruta = os.path.join(actual, nombre)
                clave = os.path.normcase(ruta)
                if clave in vistos or ctx.validador.lista_negra.motivo_nombre(nombre):
                    continue
                puntos = _puntaje(consulta, nombre)
                if puntos >= 0.6:
                    vistos.add(clave)
                    es_dir = nombre in subcarpetas
                    encontrados.append((puntos, ruta + ("\\" if es_dir else "")))
            if revisados >= lim.max_archivos_revisados or time.monotonic() > fin:
                interrumpido = True
                break
        if interrumpido:
            break

    encontrados.sort(key=lambda x: (-x[0], len(x[1])))
    maximo = int(lim.max_resultados_busqueda)
    lista = [r for _, r in encontrados[:maximo]]
    resumen = f"Encontré {len(encontrados)} coincidencias para '{args['nombre_aproximado']}'."
    if len(encontrados) > maximo:
        resumen += f" Se muestran las {maximo} mejores."
    if interrumpido:
        resumen += " La búsqueda se detuvo por límite de tiempo o de archivos revisados; puede haber más."
    if not lista:
        return exito(resumen + " No hay resultados.")
    return exito(resumen, "\n".join(lista), origen="búsqueda de archivos")


# ---------------------------------------------------------------------- #
def _leer_texto(ruta: str, max_bytes: int) -> str:
    with open(ruta, "rb") as f:
        datos = f.read(max_bytes)
    for codificacion in ("utf-8-sig", "cp1252"):
        try:
            return datos.decode(codificacion)
        except UnicodeDecodeError:
            continue
    return datos.decode("utf-8", errors="replace")


def _leer_docx(ruta: str, max_chars: int) -> str:
    import docx  # python-docx

    documento = docx.Document(ruta)
    partes, total = [], 0
    for p in documento.paragraphs:
        if p.text.strip():
            partes.append(p.text)
            total += len(p.text)
            if total > max_chars:
                return "\n".join(partes)
    for tabla in documento.tables:
        for fila in tabla.rows:
            linea = " | ".join(c.text.strip() for c in fila.cells)
            partes.append(linea)
            total += len(linea)
            if total > max_chars:
                return "\n".join(partes)
    return "\n".join(partes)


def _leer_xlsx(ruta: str, max_chars: int) -> str:
    import openpyxl

    libro = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    try:
        partes, total = [], 0
        for hoja in libro.worksheets:
            partes.append(f"--- Hoja: {hoja.title} ---")
            for fila in hoja.iter_rows(values_only=True):
                if all(v is None for v in fila):
                    continue
                linea = " | ".join("" if v is None else str(v) for v in fila)
                partes.append(linea)
                total += len(linea)
                if total > max_chars:
                    return "\n".join(partes)
        return "\n".join(partes)
    finally:
        libro.close()


def _leer_pdf(ruta: str, max_chars: int) -> str:
    from pypdf import PdfReader

    lector = PdfReader(ruta)
    if lector.is_encrypted:
        try:
            lector.decrypt("")
        except Exception:  # noqa: BLE001
            raise ErrorSeguridad("El PDF está protegido con contraseña.")
    partes, total = [], 0
    for i, pagina in enumerate(lector.pages, start=1):
        texto = pagina.extract_text() or ""
        partes.append(f"--- Página {i} ---\n{texto}")
        total += len(texto)
        if total > max_chars:
            break
    return "\n".join(partes)


def leer_archivo(ctx: Contexto, args: dict) -> Resultado:
    ruta = ctx.validador.validar(args["ruta"], tipo="archivo")
    ext = os.path.splitext(ruta)[1].lower()
    if ext not in EXTENSIONES_LEIBLES:
        return error(
            f"No puedo leer archivos {ext or 'sin extensión'}. Formatos permitidos: "
            + ", ".join(sorted(EXTENSIONES_LEIBLES))
        )
    lim = ctx.config.limites
    max_bytes = int(lim.max_tamano_archivo_mb * 1024 * 1024)
    tam = os.path.getsize(ruta)
    if tam > max_bytes:
        return error(
            f"El archivo pesa {_tamano_legible(tam)} y el máximo permitido es "
            f"{lim.max_tamano_archivo_mb} MB."
        )
    max_chars = int(lim.max_caracteres_lectura)

    if ext in EXTENSIONES_TEXTO:
        texto = _leer_texto(ruta, max_bytes)
    elif ext == ".docx":
        texto = _leer_docx(ruta, max_chars)
    elif ext == ".xlsx":
        texto = _leer_xlsx(ruta, max_chars)
    else:
        texto = _leer_pdf(ruta, max_chars)

    texto = texto.strip()
    recortado = len(texto) > max_chars
    texto = texto[:max_chars]
    resumen = f"Contenido de {ruta} ({_tamano_legible(tam)})."
    if recortado:
        resumen += f" Se muestran solo los primeros {max_chars} caracteres."
    if not texto:
        return exito(resumen + " El archivo no tiene texto extraíble.")
    return exito(resumen, texto, origen=ruta)


# ---------------------------------------------------------------------- #
def _validar_para_abrir(ctx: Contexto, ruta: str) -> str:
    real = ctx.validador.validar(ruta, tipo="archivo")
    if es_ejecutable(real, ctx.config.extensiones_ejecutables_extra):
        raise ErrorSeguridad(
            "Por seguridad no abro archivos ejecutables, scripts, accesos directos ni instaladores."
        )
    return real


def abrir_archivo(ctx: Contexto, args: dict) -> Resultado:
    real = _validar_para_abrir(ctx, args["ruta"])
    os.startfile(real)  # noqa: S606 - ruta validada y extensión no ejecutable
    return exito(f"Abrí {os.path.basename(real)} en su programa predeterminado.")


def _confirmar_abrir_archivo(ctx: Contexto, args: dict) -> str | None:
    try:
        real = _validar_para_abrir(ctx, args["ruta"])
    except ErrorSeguridad:
        return None  # la ejecución devolverá el error
    if ctx.validador.esta_en_descargas(real):
        return f"El archivo '{os.path.basename(real)}' está en Descargas. ¿Quieres abrirlo?"
    return None


def abrir_carpeta(ctx: Contexto, args: dict) -> Resultado:
    real = ctx.validador.validar(args["ruta"], tipo="carpeta")
    # Carpetas con sufijo CLSID (".{...}") se comportan como objetos del shell.
    if ".{" in os.path.basename(real):
        return error("Esa carpeta tiene un nombre especial del sistema y no se abrirá.")
    os.startfile(real)  # noqa: S606
    return exito(f"Abrí la carpeta {real} en el Explorador.")


# ---------------------------------------------------------------------- #
def herramientas() -> list[Herramienta]:
    ruta = {"type": "string", "maxLength": 1000}
    return [
        Herramienta(
            "listar_carpeta",
            "Lista archivos y subcarpetas de una carpeta permitida.",
            esquema_objeto({"ruta": ruta}, ["ruta"]),
            listar_carpeta,
        ),
        Herramienta(
            "buscar_archivos",
            "Busca archivos o carpetas por nombre aproximado dentro de las carpetas permitidas.",
            esquema_objeto({
                "nombre_aproximado": {"type": "string", "maxLength": 200},
                "carpeta_opcional": {"type": "string", "maxLength": 1000},
            }, ["nombre_aproximado"]),
            buscar_archivos,
        ),
        Herramienta(
            "leer_archivo",
            "Lee el texto de un archivo .txt, .md, .csv, .json, .log, .docx, .xlsx o .pdf.",
            esquema_objeto({"ruta": ruta}, ["ruta"]),
            leer_archivo,
            lee_contenido=True,
        ),
        Herramienta(
            "abrir_archivo",
            "Abre un archivo (no ejecutable) en su programa predeterminado.",
            esquema_objeto({"ruta": ruta}, ["ruta"]),
            abrir_archivo,
            confirmacion=_confirmar_abrir_archivo,
        ),
        Herramienta(
            "abrir_carpeta",
            "Abre una carpeta permitida en el Explorador de Windows.",
            esquema_objeto({"ruta": ruta}, ["ruta"]),
            abrir_carpeta,
        ),
    ]
